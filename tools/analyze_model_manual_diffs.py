from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill


ROOT = Path("C:/Users/PC/Desktop/数据标注")
V2_ROOT = ROOT / "ShuJuBiaoZhu-v2-sandbox"

INPUTS = {
    "标准层": Path(
        r"D:/Richang/Wechat/Liaotianjilu/xwechat_files/wxid_fa0qzgt1i72p22_c39a/temp/RWTemp/2026-07/5401ba771ad359cd6e0b979d35e16002/标准层标注结果.xlsx"
    ),
    "试验层": Path(
        r"D:/Richang/Wechat/Liaotianjilu/xwechat_files/wxid_fa0qzgt1i72p22_c39a/temp/RWTemp/2026-07/5401ba771ad359cd6e0b979d35e16002/试验层标注结果.xlsx"
    ),
}

RULES_FILE = V2_ROOT / "data/input/rules.xlsx"
OUT_DIR = V2_ROOT / "data/manual_diff_analysis"
OUT_XLSX = OUT_DIR / "qwen_deepseek_manual_diff_analysis.xlsx"

COL_TRIAL_REG = "试验注册号"
COL_TRIAL_ID = "试验标识"
COL_STANDARD_NO = "标准编号"
COL_RULE_TYPE = "规则标识"
COL_RULE_TEXT = "标准内容"
COL_PATIENT_ID = "患者编号"
COL_QWEN = "标注结果-QWen"
COL_DEEPSEEK = "标注结果_deepseek"
COL_GOLD = "标注结果_gold"
COL_MANUAL = "标注结果-人工.1"


def norm_label(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    mapping = {
        "符合": "符合",
        "是": "符合",
        "Y": "符合",
        "y": "符合",
        "yes": "符合",
        "不符合": "不符合",
        "否": "不符合",
        "N": "不符合",
        "n": "不符合",
        "no": "不符合",
        "未知": "未知",
        "不确定": "未知",
        "无法判断": "未知",
        "U": "未知",
        "u": "未知",
        "unknown": "未知",
    }
    return mapping.get(text, mapping.get(text.lower(), text))


def load_cancer_map() -> pd.DataFrame:
    rules = pd.read_excel(RULES_FILE, dtype=str)
    cols = [COL_TRIAL_ID, "癌症类型", "癌症分类"]
    cancer = rules[cols].drop_duplicates(subset=[COL_TRIAL_ID]).copy()
    return cancer


def load_result(name: str, path: Path, cancer_map: pd.DataFrame) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name="Sheet1", dtype=str)
    df.insert(0, "数据层级", name)
    for col in [COL_QWEN, COL_DEEPSEEK, COL_GOLD, COL_MANUAL]:
        df[f"{col}_norm"] = df[col].map(norm_label)

    df["人工非空"] = df[f"{COL_MANUAL}_norm"] != ""
    df["QWen是否不同于人工"] = df["人工非空"] & (df[f"{COL_QWEN}_norm"] != df[f"{COL_MANUAL}_norm"])
    df["DeepSeek是否不同于人工"] = df["人工非空"] & (df[f"{COL_DEEPSEEK}_norm"] != df[f"{COL_MANUAL}_norm"])
    df["Gold是否不同于人工"] = df["人工非空"] & (df[f"{COL_GOLD}_norm"] != df[f"{COL_MANUAL}_norm"])
    df["QWen或DeepSeek不同于人工"] = df["QWen是否不同于人工"] | df["DeepSeek是否不同于人工"]

    merged = df.merge(cancer_map, on=COL_TRIAL_ID, how="left")
    merged["癌症类型"] = merged["癌症类型"].fillna("未匹配")
    merged["癌症分类"] = merged["癌症分类"].fillna("未匹配")
    return merged


def rate(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def summarize_overall(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for level, group in df.groupby("数据层级", dropna=False):
        manual = group[group["人工非空"]]
        total = len(group)
        manual_total = len(manual)
        qwen_diff = int(manual["QWen是否不同于人工"].sum())
        deepseek_diff = int(manual["DeepSeek是否不同于人工"].sum())
        gold_diff = int(manual["Gold是否不同于人工"].sum())
        either_diff = int(manual["QWen或DeepSeek不同于人工"].sum())
        rows.append(
            {
                "数据层级": level,
                "总行数": total,
                "人工非空行数": manual_total,
                "QWen不一致数": qwen_diff,
                "QWen一致率": rate(manual_total - qwen_diff, manual_total),
                "DeepSeek不一致数": deepseek_diff,
                "DeepSeek一致率": rate(manual_total - deepseek_diff, manual_total),
                "Gold不一致数": gold_diff,
                "Gold一致率": rate(manual_total - gold_diff, manual_total),
                "QWen或DeepSeek不一致样本数": either_diff,
                "QWen或DeepSeek不一致样本占人工非空比例": rate(either_diff, manual_total),
            }
        )
    return pd.DataFrame(rows)


def grouped_summary(df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    rows = []
    manual = df[df["人工非空"]].copy()
    for keys, group in manual.groupby(group_cols, dropna=False):
        if not isinstance(keys, tuple):
            keys = (keys,)
        total = len(group)
        qwen_diff = int(group["QWen是否不同于人工"].sum())
        deepseek_diff = int(group["DeepSeek是否不同于人工"].sum())
        gold_diff = int(group["Gold是否不同于人工"].sum())
        either_diff = int(group["QWen或DeepSeek不同于人工"].sum())
        row = {col: key for col, key in zip(group_cols, keys)}
        row.update(
            {
                "人工非空行数": total,
                "QWen不一致数": qwen_diff,
                "QWen一致率": rate(total - qwen_diff, total),
                "DeepSeek不一致数": deepseek_diff,
                "DeepSeek一致率": rate(total - deepseek_diff, total),
                "Gold不一致数": gold_diff,
                "Gold一致率": rate(total - gold_diff, total),
                "QWen或DeepSeek不一致样本数": either_diff,
                "QWen或DeepSeek不一致率": rate(either_diff, total),
            }
        )
        rows.append(row)
    out = pd.DataFrame(rows)
    if not out.empty:
        out = out.sort_values(["QWen或DeepSeek不一致样本数", "QWen或DeepSeek不一致率"], ascending=[False, False])
    return out


def transition_summary(df: pd.DataFrame) -> pd.DataFrame:
    manual = df[df["人工非空"]].copy()
    rows = []
    for model_name, model_col, diff_col in [
        ("QWen", f"{COL_QWEN}_norm", "QWen是否不同于人工"),
        ("DeepSeek", f"{COL_DEEPSEEK}_norm", "DeepSeek是否不同于人工"),
        ("Gold", f"{COL_GOLD}_norm", "Gold是否不同于人工"),
    ]:
        bad = manual[manual[diff_col]]
        group_cols = ["数据层级", COL_RULE_TYPE, f"{COL_MANUAL}_norm", model_col]
        for keys, group in bad.groupby(group_cols, dropna=False):
            rows.append(
                {
                    "模型": model_name,
                    "数据层级": keys[0],
                    "规则类型": keys[1],
                    "人工标签": keys[2],
                    "模型标签": keys[3],
                    "不一致数": len(group),
                }
            )
    out = pd.DataFrame(rows)
    if not out.empty:
        out = out.sort_values(["模型", "数据层级", "不一致数"], ascending=[True, True, False])
    return out


def autosize_and_style(writer: pd.ExcelWriter) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for col in ws.columns:
            max_len = 0
            letter = col[0].column_letter
            for cell in col[:200]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    cancer_map = load_cancer_map()
    frames = [load_result(name, path, cancer_map) for name, path in INPUTS.items()]
    all_df = pd.concat(frames, ignore_index=True)

    diff_cols = [
        "数据层级",
        COL_TRIAL_REG,
        COL_TRIAL_ID,
        "癌症类型",
        "癌症分类",
        COL_STANDARD_NO,
        COL_RULE_TYPE,
        COL_RULE_TEXT,
        COL_PATIENT_ID,
        COL_QWEN,
        COL_DEEPSEEK,
        COL_GOLD,
        COL_MANUAL,
        "QWen是否不同于人工",
        "DeepSeek是否不同于人工",
        "Gold是否不同于人工",
        "QWen或DeepSeek不同于人工",
        "匹配解释",
        "匹配解释_deepseek",
        "匹配解释_gold",
        "参考原始病历信息",
        "证据来源",
        "置信度",
        "规则后处理",
    ]
    diff_rows = all_df[all_df["人工非空"] & all_df["QWen或DeepSeek不同于人工"]].copy()
    diff_rows = diff_rows[[col for col in diff_cols if col in diff_rows.columns]]

    overview = summarize_overall(all_df)
    by_rule_type = grouped_summary(all_df, ["数据层级", COL_RULE_TYPE])
    by_label = grouped_summary(all_df, ["数据层级", f"{COL_MANUAL}_norm"])
    by_trial = grouped_summary(all_df, ["数据层级", COL_TRIAL_ID, COL_TRIAL_REG, "癌症类型", "癌症分类"])
    by_cancer = grouped_summary(all_df, ["数据层级", "癌症类型", "癌症分类"])
    by_standard = grouped_summary(all_df, ["数据层级", COL_TRIAL_ID, COL_STANDARD_NO, COL_RULE_TYPE, "癌症类型", "癌症分类"])
    transitions = transition_summary(all_df)

    notes = pd.DataFrame(
        [
            {"说明项": "差异样本定义", "内容": "人工标注结果非空，且 QWen 或 DeepSeek 任一模型标签不同于人工标签。"},
            {"说明项": "人工口径", "内容": "涉及人工的一致率只统计人工列非空的行。"},
            {"说明项": "Gold处理", "内容": "Gold不参与差异样本筛选，但保留在样本和汇总中作为参考。"},
            {"说明项": "癌种来源", "内容": f"按试验标识从规则表补充：{RULES_FILE}"},
            {"说明项": "输出文件", "内容": str(OUT_XLSX)},
        ]
    )

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        notes.to_excel(writer, sheet_name="说明", index=False)
        overview.to_excel(writer, sheet_name="总览", index=False)
        diff_rows[diff_rows["数据层级"] == "标准层"].to_excel(writer, sheet_name="标准层_差异样本", index=False)
        diff_rows[diff_rows["数据层级"] == "试验层"].to_excel(writer, sheet_name="试验层_差异样本", index=False)
        by_rule_type.to_excel(writer, sheet_name="按规则类型统计", index=False)
        by_label.to_excel(writer, sheet_name="按人工标签统计", index=False)
        transitions.to_excel(writer, sheet_name="按标签转移统计", index=False)
        by_trial.to_excel(writer, sheet_name="按试验统计", index=False)
        by_cancer.to_excel(writer, sheet_name="按癌种统计", index=False)
        by_standard.to_excel(writer, sheet_name="按标准统计", index=False)
        autosize_and_style(writer)

    print(f"output={OUT_XLSX}")
    print(overview.to_string(index=False))
    print(f"diff_rows={len(diff_rows)}")


if __name__ == "__main__":
    main()
