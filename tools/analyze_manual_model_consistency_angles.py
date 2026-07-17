from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill


ROOT = Path("C:/Users/PC/Desktop/数据标注")
V2_ROOT = ROOT / "ShuJuBiaoZhu-v2-sandbox"

INPUTS = {
    "标准层": Path(
        r"D:/Richang/Wechat/Liaotianjilu/xwechat_files/wxid_fa0qzgt1i72p22_c39a/temp/RWTemp/2026-07/5401ba771ad359cd6e0b979d35e16002/标准层标注结果.xlsx"
    ),
    "试验层": Path(
        r"D:/Richang/Wechat/Liaotianjilu/xwechat_files/wxid_fa0qzgt1i72p22_c39a/temp/RWTemp/2026-07/5401ba771ad359cd6e0b979d35e16002/试验层标注结果.xlsx"
    ),
}

RULES_FILE = V2_ROOT / "data/input/rules.xlsx"
OUT_DIR = V2_ROOT / "data/manual_diff_analysis"
OUT_XLSX = OUT_DIR / "manual_qwen_deepseek_consistency_angles.xlsx"

COL_TRIAL_REG = "试验注册号"
COL_TRIAL_ID = "试验标识"
COL_STANDARD_NO = "标准编号"
COL_RULE_TYPE = "规则标识"
COL_RULE_TEXT = "标准内容"
COL_PATIENT_ID = "患者编号"
COL_QWEN = "标注结果-QWen"
COL_DEEPSEEK = "标注结果_deepseek"
COL_MANUAL = "标注结果-人工.1"


def norm_label(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    mapping = {
        "符合": "符合",
        "是": "符合",
        "yes": "符合",
        "y": "符合",
        "不符合": "不符合",
        "否": "不符合",
        "no": "不符合",
        "n": "不符合",
        "未知": "未知",
        "不确定": "未知",
        "无法判断": "未知",
        "unknown": "未知",
        "u": "未知",
    }
    return mapping.get(text, mapping.get(text.lower(), text))


def load_cancer_map() -> pd.DataFrame:
    rules = pd.read_excel(RULES_FILE, dtype=str)
    return rules[[COL_TRIAL_ID, "癌症类型", "癌症分类"]].drop_duplicates(subset=[COL_TRIAL_ID])


def load_all() -> pd.DataFrame:
    cancer = load_cancer_map()
    frames = []
    for level, path in INPUTS.items():
        df = pd.read_excel(path, sheet_name="Sheet1", dtype=str)
        df.insert(0, "数据层级", level)
        df["人工_norm"] = df[COL_MANUAL].map(norm_label)
        df["QWen_norm"] = df[COL_QWEN].map(norm_label)
        df["DeepSeek_norm"] = df[COL_DEEPSEEK].map(norm_label)
        df = df.merge(cancer, on=COL_TRIAL_ID, how="left")
        df["癌症类型"] = df["癌症类型"].fillna("未匹配")
        df["癌症分类"] = df["癌症分类"].fillna("未匹配")
        frames.append(df)
    all_df = pd.concat(frames, ignore_index=True)
    return all_df[all_df["人工_norm"] != ""].copy()


def add_model_rows(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for model, col in [("QWen", "QWen_norm"), ("DeepSeek", "DeepSeek_norm")]:
        part = df.copy()
        part["模型"] = model
        part["模型标签"] = part[col]
        part["是否一致"] = part["模型标签"] == part["人工_norm"]
        rows.append(part)
    return pd.concat(rows, ignore_index=True)


def summarize(group: pd.core.groupby.DataFrameGroupBy) -> pd.DataFrame:
    out = group["是否一致"].agg(可比行数="count", 一致数="sum", 一致率="mean").reset_index()
    out["不一致数"] = out["可比行数"] - out["一致数"]
    return out


def low_rate_summary(model_df: pd.DataFrame, group_cols: list[str], threshold: float = 0.95) -> pd.DataFrame:
    out = summarize(model_df.groupby(["模型", *group_cols], dropna=False))
    out = out[out["一致率"] < threshold].copy()
    return out.sort_values(["一致率", "不一致数", "可比行数"], ascending=[True, False, False])


def rules_low_rate(model_df: pd.DataFrame) -> pd.DataFrame:
    group_cols = [
        "数据层级",
        COL_TRIAL_ID,
        COL_TRIAL_REG,
        "癌症类型",
        "癌症分类",
        COL_STANDARD_NO,
        COL_RULE_TYPE,
        COL_RULE_TEXT,
    ]
    out = low_rate_summary(model_df, group_cols)
    cols = [
        "模型",
        *group_cols,
        "可比行数",
        "一致数",
        "不一致数",
        "一致率",
    ]
    return out[cols]


def transition_by_manual_label(model_df: pd.DataFrame) -> pd.DataFrame:
    return summarize(model_df.groupby(["模型", "人工_norm"], dropna=False)).sort_values(
        ["模型", "一致率", "可比行数"], ascending=[True, True, False]
    )


def style_workbook(writer: pd.ExcelWriter) -> None:
    fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
        for column in ws.columns:
            width = 10
            letter = column[0].column_letter
            for cell in column[:200]:
                width = max(width, len(str(cell.value or "")) + 2)
            ws.column_dimensions[letter].width = min(width, 48)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    manual_df = load_all()
    model_df = add_model_rows(manual_df)

    overview = summarize(model_df.groupby(["模型"], dropna=False))
    by_label = transition_by_manual_label(model_df)
    by_level = summarize(model_df.groupby(["模型", "数据层级"], dropna=False)).sort_values(
        ["模型", "一致率"], ascending=[True, True]
    )
    low_rules = rules_low_rate(model_df)
    by_cancer = summarize(model_df.groupby(["模型", "癌症类型", "癌症分类"], dropna=False)).sort_values(
        ["模型", "一致率", "不一致数"], ascending=[True, True, False]
    )
    low_cancer = by_cancer[by_cancer["一致率"] < 0.95].copy()
    by_rule_type = summarize(model_df.groupby(["模型", COL_RULE_TYPE], dropna=False)).sort_values(
        ["模型", "一致率"], ascending=[True, True]
    )

    notes = pd.DataFrame(
        [
            {"问题": "统计对象", "说明": "只统计人工标注非空行；只比较 QWen/DeepSeek 与人工，不纳入 Gold。"},
            {"问题": "低一致率阈值", "说明": "规则和癌种低一致率阈值为 <95%。"},
            {"问题": "癌种来源", "说明": f"按试验标识从规则表补充：{RULES_FILE}"},
            {"问题": "输出文件", "说明": str(OUT_XLSX)},
        ]
    )

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        notes.to_excel(writer, sheet_name="说明", index=False)
        overview.to_excel(writer, sheet_name="总览", index=False)
        by_label.to_excel(writer, sheet_name="1_按人工标签", index=False)
        by_level.to_excel(writer, sheet_name="2_实验层vs标准层", index=False)
        low_rules.to_excel(writer, sheet_name="3_低于95规则", index=False)
        low_cancer.to_excel(writer, sheet_name="4_低于95癌种", index=False)
        by_cancer.to_excel(writer, sheet_name="4_全部癌种", index=False)
        by_rule_type.to_excel(writer, sheet_name="5_入选排除", index=False)
        style_workbook(writer)

    print(f"output={OUT_XLSX}")
    print("overview")
    print(overview.to_string(index=False))
    print("by_label")
    print(by_label.to_string(index=False))
    print("by_level")
    print(by_level.to_string(index=False))
    print("by_rule_type")
    print(by_rule_type.to_string(index=False))
    print("low_rules_count", len(low_rules))
    print("low_cancer_count", len(low_cancer))


if __name__ == "__main__":
    main()
