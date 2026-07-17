from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl


@dataclass(frozen=True)
class TrialRule:
    trial_id: str
    trial_register_id: str
    protocol_no: str
    cancer_type: str
    cancer_category: str
    standard_no: str
    rule_text: str
    rule_type: str
    source_sheet: str = ""
    source_row: int = 0


HEADER_ALIASES = {
    "trial_id": ["试验标识"],
    "trial_register_id": ["试验注册号"],
    "protocol_no": ["试验方案编号"],
    "cancer_type": ["癌症类型"],
    "cancer_category": ["癌症分类"],
    "standard_no": ["标准编号"],
    "rule_text": ["规则"],
    "rule_type": ["规则标识"],
}


def _header_map(headers: list[Any], sheet_name: str) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, header in enumerate(headers):
        if header is None:
            continue
        name = str(header).strip()
        for field, aliases in HEADER_ALIASES.items():
            if name in aliases and field not in mapping:
                mapping[field] = idx
    missing = [field for field in HEADER_ALIASES if field not in mapping]
    if missing:
        raise ValueError(f"Rules sheet {sheet_name!r} missing columns: {missing}")
    return mapping


def _cell(row: tuple[Any, ...], mapping: dict[str, int], name: str) -> str:
    idx = mapping[name]
    if idx >= len(row):
        return ""
    return str(row[idx] or "").strip()


def load_rules(path: Path, limit: int | None = None) -> list[TrialRule]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rules: list[TrialRule] = []
    seen_keys: set[tuple[str, str]] = set()
    try:
        for ws in wb.worksheets:
            if ws.max_row < 2:
                continue
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            mapping = _header_map(list(header_row), ws.title)
            for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row):
                    continue
                rule_text = _cell(row, mapping, "rule_text")
                if not rule_text:
                    continue
                trial_id = _cell(row, mapping, "trial_id")
                standard_no = _cell(row, mapping, "standard_no")
                key = (trial_id, standard_no)
                if key in seen_keys:
                    raise ValueError(
                        f"Duplicate rule key in {path}: trial_id={trial_id!r}, standard_no={standard_no!r}"
                    )
                seen_keys.add(key)
                rules.append(
                    TrialRule(
                        trial_id=trial_id,
                        trial_register_id=_cell(row, mapping, "trial_register_id"),
                        protocol_no=_cell(row, mapping, "protocol_no"),
                        cancer_type=_cell(row, mapping, "cancer_type"),
                        cancer_category=_cell(row, mapping, "cancer_category"),
                        standard_no=standard_no,
                        rule_text=rule_text,
                        rule_type=_cell(row, mapping, "rule_type"),
                        source_sheet=ws.title,
                        source_row=row_number,
                    )
                )
                if limit and len(rules) >= limit:
                    return rules
        return rules
    finally:
        wb.close()
